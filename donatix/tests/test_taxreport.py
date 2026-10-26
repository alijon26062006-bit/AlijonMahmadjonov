"""Учёт для ИП: точная закупка от поставщика, заказы в минусе, файл Excel с налогом."""
import io
from datetime import datetime, timezone
from decimal import Decimal

from conftest import make_client, web_login
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from donatix import finance, orders, taxreport
from donatix.suppliers.base import charged_of
from test_finance import _order, _pay

UTC = timezone.utc


def _day(conn):
    return finance.window(conn, datetime(2026, 9, 27, 8, 0, tzinfo=UTC), days_back=1)


def test_excel_has_turnover_tax_orders_and_losses(config, conn):
    uid, _ = make_client(conn)
    taxreport.set_tax_percent(conn, "6")
    _pay(conn, uid, 100, "1100.00", "TJS", "2026-09-26T08:00:00")
    _order(conn, uid, 1, 60, 50, "2026-09-26T09:00:00")
    _order(conn, uid, 2, 40, 45, "2026-09-26T15:00:00")      # продали дешевле закупки — минус
    start, end = _day(conn)
    wb = load_workbook(io.BytesIO(taxreport.build(conn, config, start, end)))
    assert wb.sheetnames == ["Итог по дням", "Заказы", "Поступления"]
    day = [c.value for c in wb["Итог по дням"][4]]
    assert day[0] == "26.09.2026" and day[1] == 1100.0           # оборот — как перевели, в сомони
    assert day[9] == 66.0                                        # налог 6 % от оборота
    rows = list(wb["Заказы"].iter_rows(min_row=4, values_only=True))
    loss = next(r for r in rows if r and r[1] == "dx-f2")
    assert loss[8] == -5.0 and loss[7] == -11.1                  # прибыль $-5, наценка −11.1 %
    assert any(r and r[0] and "В минусе: 1" in str(r[0]) for r in rows)
    assert wb["Поступления"]["E4"].value == 1100.0
    assert taxreport.filename(config, start, end).endswith("2026-09-26.xlsx")


def test_tax_percent_validation(conn):
    assert taxreport.tax_percent(conn) == 0
    assert taxreport.set_tax_percent(conn, "6,5%") == 6.5
    for bad in ("abc", "150", "-1"):
        try:
            taxreport.set_tax_percent(conn, bad)
            raise AssertionError(bad)
        except ValueError:
            pass


def test_real_supplier_cost_replaces_catalog(conn):
    uid, _ = make_client(conn)
    _order(conn, uid, 9, 10, 8, "2026-09-26T09:00:00")
    o = conn.execute("SELECT * FROM orders WHERE public_id = 'dx-f9'").fetchone()
    assert charged_of({"id": 1, "total_price": "8.4000"}) == Decimal("8.4") and charged_of({"status": "ok"}) is None
    assert not orders.real_cost(conn, o, 50)                     # в 6 раз дороже каталога — это не цена, не верим
    assert orders.real_cost(conn, o, "8.4")
    got = conn.execute("SELECT cost_micro, cost_source FROM orders WHERE id = ?", (o["id"],)).fetchone()
    assert (got["cost_micro"], got["cost_source"]) == (84_000, "supplier")


def test_admin_money_page_shows_losses_and_download(app, config, conn):
    uid, _ = make_client(conn)
    now = datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%S")
    _order(conn, uid, 3, 40, 45, now)
    c = TestClient(app)
    web_login(c, "admin@example.com", "adminpass123")
    page = c.get("/admin/money?period=today").text
    assert "В минусе" in page and "Файл для ИП" in page and "Налог, % от оборота" in page
    r = c.get("/admin/money/export?period=today")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in r.headers["content-disposition"]
    anon = TestClient(app).get("/admin/money/export?period=today", follow_redirects=False)
    assert anon.status_code in (302, 303, 401, 403)                # без входа файл не отдаём


def test_daily_report_sends_excel(config, conn, monkeypatch):
    sent = []
    monkeypatch.setattr("donatix.worker.notify_admin", lambda *a, **k: None)
    monkeypatch.setattr("donatix.worker.notify_admin_file",
                        lambda cfg, caption, buttons, path, photo: sent.append((path.name, path.read_bytes()[:2])))
    assert finance.maybe_send(conn, config, datetime(2026, 9, 27, 8, 0, tzinfo=UTC))
    assert sent and sent[0][0].endswith(".xlsx") and sent[0][1] == b"PK"
