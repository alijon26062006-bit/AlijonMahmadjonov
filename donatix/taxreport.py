"""Отчёт для ИП и налогов — один файл Excel за день, месяц или любой период.

Лист «Итог по дням»: сколько пришло от клиентов (оборот), сколько продали, во что обошлась закупка
у поставщика, прибыль и налог с прибыли по проценту из настроек — по каждому дню и итого.
Лист «Заказы»: каждый заказ — продажа, закупка (точно от поставщика или по каталогу), наценка в %,
прибыль; убыточные выделены красным. Лист «Поступления»: каждое пополнение от клиента.

Суммы в сомони: продажи — по курсу в момент заказа, поступления в сомони — как перевели,
остальные — по текущему курсу. Налог считается от чистой прибыли (продажи − закупка − рефералы и D-коин);
день в минусе — налог 0. Процент задаёт владелец.
"""

from __future__ import annotations

import io
import sqlite3
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any

from . import db, finance, timez
from .config import Config

TAX_KEY = "tax.percent"


def tax_percent(conn: sqlite3.Connection) -> Decimal:
    try:
        value = Decimal((db.get_setting(conn, TAX_KEY) or "0").replace(",", "."))
    except InvalidOperation:
        return Decimal(0)
    return value if 0 <= value <= 100 else Decimal(0)


def set_tax_percent(conn: sqlite3.Connection, raw: str) -> Decimal:
    try:
        value = Decimal((raw or "0").strip().replace(",", ".").rstrip("%"))
    except InvalidOperation:
        raise ValueError("Налог — числом, например 6") from None
    if not 0 <= value <= 100:
        raise ValueError("Налог — от 0 до 100 %")
    db.set_setting(conn, TAX_KEY, str(value.normalize()))
    return value


def _tjs(micro: int, rate: Decimal) -> float:
    return float((Decimal(micro) / 10_000 * rate).quantize(Decimal("0.01")))


def _usd(micro: int) -> float:
    return float((Decimal(micro) / 10_000).quantize(Decimal("0.0001")))


def _rate(raw: Any, fallback: Decimal) -> Decimal:
    try:
        value = Decimal(str(raw))
        return value if value > 0 else fallback
    except (InvalidOperation, TypeError):
        return fallback


def orders(conn: sqlite3.Connection, a: str, b: str, rate: Decimal, zone: str) -> list[dict[str, Any]]:
    """Выполненные заказы за [a, b) — с прибылью и наценкой. Плюс боты для гильдии и номера."""
    out = []
    for o in conn.execute(
            "SELECT o.*, u.login FROM orders o LEFT JOIN users u ON u.id = o.user_id WHERE o.status = 'completed' "
            "AND COALESCE(o.completed_at, o.created_at) >= ? AND COALESCE(o.completed_at, o.created_at) < ? "
            "ORDER BY COALESCE(o.completed_at, o.created_at)", (a, b)).fetchall():
        keys = o.keys()
        r = _rate(o["tjs_rate"] if "tjs_rate" in keys else None, rate)
        sale, cost = int(o["total_micro"]), int(o["cost_micro"])
        exact = "cost_source" in keys and o["cost_source"] == "supplier"
        out.append({"time": timez.local(o["completed_at"] or o["created_at"], zone), "id": o["public_id"],
                    "product": o["product_name"], "client": o["login"] or "", "sale": sale, "cost": cost,
                    "source": "точно (поставщик)" if exact else "по каталогу", "rate": r})
    for g in _safe(conn, "SELECT g.*, u.login FROM glory_orders g LEFT JOIN users u ON u.id = g.user_id "
                         "WHERE g.refunded = 0 AND g.status != 'failed' AND g.created_at >= ? AND g.created_at < ?",
                   (a, b)):
        cost = int(Decimal(str(g["cost_usd"] or 0)) * 10_000) if str(g["cost_usd"] or "").strip() else 0
        out.append({"time": timez.local(g["created_at"], zone), "id": f"gl-{g['id']}",
                    "product": f"Боты для гильдии, клан {g['clan_id']}", "client": g["login"] or "",
                    "sale": int(g["price_micro"]), "cost": cost, "source": "цена кредита", "rate": rate})
    return out


def payments(conn: sqlite3.Connection, a: str, b: str, rate: Decimal, zone: str) -> list[dict[str, Any]]:
    """Зачисленные пополнения за [a, b). В сомони: перевели в сомони — как есть, иначе — по курсу."""
    out = []
    for p in conn.execute(
            "SELECT p.*, u.login FROM payments p LEFT JOIN users u ON u.id = p.user_id WHERE p.status = 'paid' "
            "AND p.resolved_at >= ? AND p.resolved_at < ? ORDER BY p.resolved_at", (a, b)).fetchall():
        try:
            paid = Decimal(str(p["pay_amount"]))
        except InvalidOperation:
            paid = Decimal(0)
        tjs = paid if p["pay_currency"] == "TJS" else Decimal(p["amount_micro"]) / 10_000 * rate
        out.append({"time": timez.local(p["resolved_at"], zone), "id": p["id"], "client": p["login"] or "",
                    "method": p["method"],
                    "paid": float(paid), "currency": p["pay_currency"], "usd": p["amount_micro"],
                    "tjs": float(tjs.quantize(Decimal("0.01")))})
    return out


def _safe(conn: sqlite3.Connection, sql: str, args: tuple) -> list[sqlite3.Row]:
    try:
        return conn.execute(sql, args).fetchall()
    except sqlite3.OperationalError:   # старая база без таблицы
        return []


def build(conn: sqlite3.Connection, config: Config, start: datetime, end: datetime) -> bytes:
    """Файл .xlsx за [start, end) — местные полуночи."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from .payments import settings as pay_settings
    from .payments import title_for
    zone = timez.site_zone_name(conn)
    rate = _rate(pay_settings(conn, config)["tjs_rate"], Decimal(1))
    tax = tax_percent(conn)
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="E8EEFF")
    red, red_fill = Font(color="C42B2B", bold=True), PatternFill("solid", fgColor="FDEEEE")
    money = "#,##0.00"

    def sheet_head(ws, title: str, cols: list[tuple[str, int]]) -> None:
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([])
        ws.append([c for c, _ in cols])
        for i, (_, width) in enumerate(cols, 1):
            cell = ws.cell(row=3, column=i)
            cell.font, cell.fill = bold, head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A4"

    period = f"{start.strftime('%d.%m.%Y')} — {(end - timedelta(seconds=1)).strftime('%d.%m.%Y')}"
    wb = Workbook()

    # ── Итог по дням ──
    ws = wb.active
    ws.title = "Итог по дням"
    sheet_head(ws, f"{config.site_name}: оборот и прибыль за {period} (курс 1 $ = {rate} с.)", [
        ("Дата", 12), ("Поступило от клиентов (оборот), с.", 20), ("Пополнений", 12), ("Продано заказов", 12),
        ("Продажи, с.", 14), ("Закупка у поставщика, с.", 18), ("Прибыль с продаж, с.", 16),
        ("Рефералы и D-коин, с.", 16), ("Чистая прибыль, с.", 16), (f"Налог {tax}% с прибыли, с.", 18)])
    totals = [0.0] * 9
    cur = start
    while cur < end:
        nxt = min(cur + timedelta(days=1), end)
        s = finance.summary(conn, config, cur, nxt)
        received = sum(p["tjs"] for p in payments(conn, finance._iso(cur), finance._iso(nxt), rate, zone))
        sold = orders(conn, finance._iso(cur), finance._iso(nxt), rate, zone)
        sales = sum(_tjs(o["sale"], o["rate"]) for o in sold)
        cost = sum(_tjs(o["cost"], o["rate"]) for o in sold)
        extra = _tjs(s["referral"] + s["dcoin"], rate)
        row = [received, s["payments"], len(sold), sales, cost, sales - cost, extra, sales - cost - extra,
               round(max(sales - cost - extra, 0) * float(tax) / 100, 2)]
        ws.append([cur.strftime("%d.%m.%Y")] + [round(v, 2) if isinstance(v, float) else v for v in row])
        if row[7] < 0:
            for c in ws[ws.max_row]:
                c.fill = red_fill
        totals = [t + v for t, v in zip(totals, row)]
        cur = nxt
    ws.append(["ИТОГО"] + [round(t, 2) if i not in (1, 2) else int(t) for i, t in enumerate(totals)])
    for c in ws[ws.max_row]:
        c.font = bold
    for row in ws.iter_rows(min_row=4, min_col=2, max_col=10):
        for c in row:
            if c.column not in (3, 4):
                c.number_format = money
    ws.append([])
    ws.append([f"Налог считается с чистой прибыли (продажи − закупка − рефералы и D-коин) по ставке {tax}%; "
               "день в минусе — налог 0. Ставку задаёт владелец в «Учёте денег»."])

    # ── Заказы ──
    ws = wb.create_sheet("Заказы")
    sheet_head(ws, f"Все выполненные заказы за {period}", [
        ("Время", 17), ("Номер", 11), ("Товар", 36), ("Клиент", 14), ("Продажа, $", 11), ("Закупка, $", 11),
        ("Закупка — откуда", 16), ("Наценка, %", 10), ("Прибыль, $", 11), ("Курс", 8), ("Продажа, с.", 12),
        ("Прибыль, с.", 12)])
    a, b = finance._iso(start), finance._iso(end)
    losses = 0
    for o in orders(conn, a, b, rate, zone):
        profit = o["sale"] - o["cost"]
        markup = round((o["sale"] - o["cost"]) * 100 / o["cost"], 1) if o["cost"] else None
        ws.append([o["time"], o["id"], o["product"], o["client"], _usd(o["sale"]), _usd(o["cost"]), o["source"],
                   markup, _usd(profit), float(o["rate"]), _tjs(o["sale"], o["rate"]), _tjs(profit, o["rate"])])
        if profit < 0:
            losses += 1
            for c in ws[ws.max_row]:
                c.font = red
    if losses:
        ws.append([])
        ws.append([f"В минусе: {losses} заказ(ов) — выделены красным"])
        ws.cell(row=ws.max_row, column=1).font = red

    # ── Поступления ──
    ws = wb.create_sheet("Поступления")
    sheet_head(ws, f"Пополнения от клиентов за {period}", [
        ("Время", 17), ("№", 8), ("Клиент", 14), ("Способ", 22), ("Перевели", 12), ("Валюта", 8),
        ("Зачислено, $", 12), ("В сомони", 12)])
    for p in payments(conn, a, b, rate, zone):
        ws.append([p["time"], p["id"], p["client"], title_for(conn, config, p["method"]), p["paid"], p["currency"],
                   _usd(p["usd"]), p["tjs"]])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(config: Config, start: datetime, end: datetime) -> str:
    last = end - timedelta(seconds=1)
    span = start.strftime("%Y-%m-%d") if start.date() == last.date() else \
        f"{start.strftime('%Y-%m-%d')}_{last.strftime('%Y-%m-%d')}"
    return f"{config.site_name.lower()}-otchet-{span}.xlsx"
