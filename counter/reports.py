"""Отчёты: Excel и PDF. Каждая запись по отдельности плюс общий итог."""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from datetime import datetime, timezone
from io import BytesIO

from .numbers import SCALE, format_amount

log = logging.getLogger(__name__)

FONT = "Report"
FONT_BOLD = "Report-Bold"
_fonts_ready = False

# Если совсем не нашлось шрифта с кириллицей — отчёт всё равно сделаем,
# просто латиницей. Лучше так, чем пустые квадраты или ошибка.
_TRANSLIT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ж": "zh", "з": "z",
    "и": "i", "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p",
    "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "h", "ц": "c", "ч": "ch",
    "ш": "sh", "щ": "sch", "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
    "ё": "e", "ў": "o", "қ": "q", "ғ": "g", "ҳ": "h",
}


def register_fonts(regular: str | None, bold: str | None) -> bool:
    """Подключить шрифт с кириллицей. False — не нашли, будет латиница."""
    global _fonts_ready
    if _fonts_ready:
        return True
    if not regular:
        return False
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        pdfmetrics.registerFont(TTFont(FONT, regular))
        pdfmetrics.registerFont(TTFont(FONT_BOLD, bold or regular))
        _fonts_ready = True
    except Exception as exc:  # noqa: BLE001
        log.warning("Шрифт не подключился (%s) — PDF будет латиницей", exc)
        return False
    return True


def _text(value: str) -> str:
    if _fonts_ready:
        return value
    return "".join(
        _TRANSLIT.get(ch.lower(), ch).upper() if ch.isupper() else _TRANSLIT.get(ch, ch)
        for ch in value
    )


@dataclass
class ReportData:
    chat_title: str
    session_title: str
    period: str
    unit: str = "шт"
    rows: list[dict] = field(default_factory=list)      # {name, count, total}
    entries: list[dict] = field(default_factory=list)   # {n, name, amount, time, note}
    generated_at: datetime = field(default_factory=lambda: datetime.now(timezone.utc))

    @property
    def total(self) -> int:
        return sum(row["total"] for row in self.rows)

    @property
    def count(self) -> int:
        return sum(row["count"] for row in self.rows)

    @property
    def file_stem(self) -> str:
        stamp = self.generated_at.strftime("%Y-%m-%d_%H-%M")
        return f"otchet_{stamp}"


def _amount(value: int) -> float | int:
    """Число для Excel: целое остаётся целым, дробное — дробным."""
    whole, rest = divmod(value, SCALE)
    return whole if rest == 0 else value / SCALE


# ── Excel ──────────────────────────────────────────────────────────────────
def build_xlsx(data: ReportData) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    book = Workbook()
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    total_fill = PatternFill("solid", fgColor="FFF2CC")

    def header(sheet, titles: list[str], widths: list[int], row: int) -> None:
        for column, (title, width) in enumerate(zip(titles, widths), start=1):
            cell = sheet.cell(row=row, column=column, value=title)
            cell.fill, cell.font, cell.border = head_fill, head_font, border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.row_dimensions[row].height = 22

    # Лист 1 — сводка
    summary = book.active
    summary.title = "Сводка"
    summary["A1"] = data.chat_title or "Счёт товара"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = f"{data.session_title} · {data.period}"
    summary["A3"] = "Отчёт составлен: " + data.generated_at.strftime("%d.%m.%Y %H:%M")
    summary["A3"].font = Font(color="808080", size=9)

    header(summary, ["№", "Человек", "Записей", f"Итого, {data.unit}", "Доля"], [6, 26, 10, 16, 10], 5)
    row_index = 6
    for number, row in enumerate(data.rows, start=1):
        share = (row["total"] / data.total) if data.total else 0
        values = [number, row["name"], row["count"], _amount(row["total"]), share]
        for column, value in enumerate(values, start=1):
            cell = summary.cell(row=row_index, column=column, value=value)
            cell.border = border
            if column == 5:
                cell.number_format = "0.0%"
            if column in (1, 3, 5):
                cell.alignment = Alignment(horizontal="center")
            if column == 4:
                cell.number_format = "# ##0.##"
                cell.font = Font(bold=True)
        row_index += 1

    for column, value in enumerate(["", "ВСЕГО", data.count, _amount(data.total), 1 if data.total else 0], start=1):
        cell = summary.cell(row=row_index, column=column, value=value)
        cell.fill, cell.border = total_fill, border
        cell.font = Font(bold=True, size=12)
        if column == 5:
            cell.number_format = "0.0%"
        if column == 4:
            cell.number_format = "# ##0.##"
    summary.freeze_panes = "A6"

    # Лист 2 — все записи по одной
    sheet = book.create_sheet("Записи")
    header(sheet, ["№", "Человек", f"Число, {data.unit}", "Время", "Как записано"], [6, 22, 14, 18, 34], 1)
    for index, entry in enumerate(data.entries, start=2):
        values = [
            entry["n"],
            entry["name"],
            _amount(entry["amount"]),
            entry["time"],
            entry.get("note", ""),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=index, column=column, value=value)
            cell.border = border
            if column == 3:
                cell.number_format = "# ##0.##"
            if column in (1, 4):
                cell.alignment = Alignment(horizontal="center")
    last = len(data.entries) + 2
    sheet.cell(row=last, column=2, value="ВСЕГО").font = Font(bold=True, size=12)
    total_cell = sheet.cell(row=last, column=3, value=_amount(data.total))
    total_cell.font = Font(bold=True, size=12)
    total_cell.number_format = "# ##0.##"
    for column in range(1, 6):
        sheet.cell(row=last, column=column).fill = total_fill
    sheet.freeze_panes = "A2"
    if data.entries:
        sheet.auto_filter.ref = f"A1:E{last - 1}"

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# ── PDF ────────────────────────────────────────────────────────────────────
def build_pdf(data: ReportData) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    font = FONT if _fonts_ready else "Helvetica"
    font_bold = FONT_BOLD if _fonts_ready else "Helvetica-Bold"

    title_style = ParagraphStyle("t", fontName=font_bold, fontSize=16, leading=20)
    sub_style = ParagraphStyle("s", fontName=font, fontSize=10, leading=14, textColor=colors.HexColor("#666666"))
    cell_style = ParagraphStyle("c", fontName=font, fontSize=9, leading=11)

    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=_text("Отчёт"),
    )

    def table_style(header_color: str = "#1F4E79") -> TableStyle:
        return TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("FONTNAME", (0, 0), (-1, 0), font_bold),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(header_color)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
            ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
            ("ALIGN", (0, 0), (0, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ])

    def total_row_style(row: int) -> list:
        return [
            ("BACKGROUND", (0, row), (-1, row), colors.HexColor("#FFF2CC")),
            ("FONTNAME", (0, row), (-1, row), font_bold),
            ("FONTSIZE", (0, row), (-1, row), 11),
        ]

    story: list = [
        Paragraph(_text(data.chat_title or "Счёт товара"), title_style),
        Paragraph(_text(f"{data.session_title} · {data.period}"), sub_style),
        Paragraph(_text("Составлен: " + data.generated_at.strftime("%d.%m.%Y %H:%M")), sub_style),
        Spacer(1, 8 * mm),
        Paragraph(_text("Итого по людям"), ParagraphStyle("h", fontName=font_bold, fontSize=12, leading=16)),
        Spacer(1, 3 * mm),
    ]

    rows = [[_text(x) for x in ["№", "Человек", "Записей", f"Итого, {data.unit}", "Доля"]]]
    for number, row in enumerate(data.rows, start=1):
        share = f"{row['total'] / data.total * 100:.1f}%" if data.total else "—"
        rows.append([
            str(number),
            Paragraph(_text(row["name"]), cell_style),
            str(row["count"]),
            format_amount(row["total"]),
            share,
        ])
    rows.append(["", Paragraph(_text("ВСЕГО"), cell_style), str(data.count), format_amount(data.total), "100%"])

    table = Table(rows, colWidths=[12 * mm, 62 * mm, 24 * mm, 45 * mm, 25 * mm], repeatRows=1)
    table.setStyle(table_style())
    table.setStyle(TableStyle(total_row_style(len(rows) - 1)))
    story.append(table)

    if data.entries:
        story.append(PageBreak())
        story.append(Paragraph(_text("Все записи по порядку"), ParagraphStyle("h2", fontName=font_bold, fontSize=12, leading=16)))
        story.append(Spacer(1, 3 * mm))
        detail = [[_text(x) for x in ["№", "Человек", f"Число, {data.unit}", "Время", "Как записано"]]]
        for entry in data.entries:
            detail.append([
                str(entry["n"]),
                Paragraph(_text(entry["name"]), cell_style),
                format_amount(entry["amount"]),
                entry["time"],
                Paragraph(_text(entry.get("note", "")), cell_style),
            ])
        detail.append(["", Paragraph(_text("ВСЕГО"), cell_style), format_amount(data.total), "", ""])
        detail_table = Table(detail, colWidths=[12 * mm, 40 * mm, 30 * mm, 26 * mm, 60 * mm], repeatRows=1)
        detail_table.setStyle(table_style("#2E6B3E"))
        detail_table.setStyle(TableStyle(total_row_style(len(detail) - 1)))
        story.append(detail_table)

    document.build(story)
    return buffer.getvalue()
