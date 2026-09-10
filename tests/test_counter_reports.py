"""Отчёты должны содержать каждую запись и правильный общий итог."""

from __future__ import annotations

from io import BytesIO

import pytest

from counter import reports
from counter.config import load_config
from counter.numbers import SCALE, format_amount


@pytest.fixture()
def data() -> reports.ReportData:
    return reports.ReportData(
        chat_title="Склад — летний товар",
        session_title="Смена 10.09.2026",
        period="с 09:00",
        rows=[
            {"name": "Иброхим", "count": 2, "total": 13000 * SCALE},
            {"name": "Азиз", "count": 1, "total": 1000 * SCALE},
        ],
        entries=[
            {"n": 1, "name": "Иброхим", "amount": 7600 * SCALE, "time": "09:12", "note": "семь шестьсот"},
            {"n": 2, "name": "Иброхим", "amount": 5400 * SCALE, "time": "09:20", "note": "пять четыреста"},
            {"n": 3, "name": "Азиз", "amount": 1000 * SCALE, "time": "09:31", "note": "1000"},
        ],
    )


def test_totals_are_summed(data):
    assert data.total == 14000 * SCALE
    assert data.count == 3
    assert format_amount(data.total) == "14 000".replace(" ", " ")


def test_xlsx_has_every_entry_and_total(data):
    from openpyxl import load_workbook

    book = load_workbook(BytesIO(reports.build_xlsx(data)))
    assert book.sheetnames == ["Сводка", "Записи"]

    sheet = book["Записи"]
    numbers = [sheet.cell(row=row, column=3).value for row in range(2, 5)]
    assert numbers == [7600, 5400, 1000]
    assert sheet.cell(row=5, column=2).value == "ВСЕГО"
    assert sheet.cell(row=5, column=3).value == 14000

    summary = book["Сводка"]
    assert summary.cell(row=6, column=2).value == "Иброхим"
    assert summary.cell(row=6, column=4).value == 13000
    assert summary.cell(row=8, column=2).value == "ВСЕГО"
    assert summary.cell(row=8, column=4).value == 14000


def test_pdf_is_readable_with_cyrillic(data):
    config = load_config(require_token=False)
    if not reports.register_fonts(config.font_path, config.font_bold_path):
        pytest.skip("в системе нет шрифта с кириллицей")

    pypdf = pytest.importorskip("pypdf")
    blob = reports.build_pdf(data)
    assert blob.startswith(b"%PDF")

    text = "".join(page.extract_text() for page in pypdf.PdfReader(BytesIO(blob)).pages)
    text = text.replace(" ", " ").replace("\xa0", " ")
    assert "Иброхим" in text and "Азиз" in text
    assert "14 000" in text          # общий итог
    assert "7 600" in text and "5 400" in text and "1 000" in text
    assert "ВСЕГО" in text


def test_pdf_survives_without_font(data, monkeypatch):
    """Нет кириллического шрифта — отчёт всё равно собирается, латиницей."""
    monkeypatch.setattr(reports, "_fonts_ready", False)
    blob = reports.build_pdf(data)
    assert blob.startswith(b"%PDF")
